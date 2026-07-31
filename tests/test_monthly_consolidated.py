from __future__ import annotations

import io
import unittest
from datetime import date, datetime, timedelta

import pandas as pd
from openpyxl import load_workbook

from hr_dashboard.hr_export import export_monthly_consolidated_workbook
from hr_dashboard.hr_incidents import (
    friendly_hr_department,
    monthly_hr_action_recommendations,
)


def incident(employee_id, name, department, day, category, incident_type):
    return {
        "employee_id": employee_id,
        "name": name,
        "department": department,
        "source_department": department,
        "date": day,
        "category": category,
        "incident_type": incident_type,
        "source_description": incident_type,
        "sheet": "Hoja 1",
    }


class MonthlyConsolidatedTests(unittest.TestCase):
    def setUp(self):
        start = date(2026, 7, 1)
        rows = [
            incident(
                "1",
                "Ana",
                "",
                start + timedelta(days=offset),
                "Incapacidad",
                "Accidente de trayecto",
            )
            for offset in range(15)
        ]
        rows.extend(
            incident(
                "2",
                "Beto",
                "OPERACIONES",
                start + timedelta(days=offset),
                "Ausencia",
                "Falta injustificada",
            )
            for offset in range(3)
        )
        rows.append(
            incident(
                "3",
                "Carla",
                "OPERACIONES",
                start,
                "Otra incidencia",
                "Permiso con goce",
            )
        )
        self.incidents = pd.DataFrame(rows)
        self.result = {
            "file_name": "julio.xls",
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "report_start": start,
            "report_end": date(2026, 7, 31),
            "incidents": self.incidents,
            "employees_in_report": 3,
        }

    def test_monthly_rules_detect_priority_cases(self):
        recommendations = monthly_hr_action_recommendations(self.incidents)

        high_priority = recommendations[recommendations["priority"] == "Alta"]
        self.assertEqual(len(high_priority), 3)
        self.assertTrue(
            high_priority["finding"].str.contains("sin departamento").any()
        )
        self.assertTrue(high_priority["finding"].str.contains("15 o más").any())
        self.assertTrue(high_priority["finding"].str.contains("3 o más").any())

    def test_monthly_export_includes_action_sheet(self):
        recommendations = monthly_hr_action_recommendations(self.incidents)
        workbook_bytes = export_monthly_consolidated_workbook(
            self.result,
            self.incidents,
            recommendations,
        )
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True)

        self.assertIn("Toma de acción", workbook.sheetnames)
        self.assertEqual(
            workbook["Dashboard"]["A1"].value,
            "Consolidado mensual de incidencias RH",
        )
        self.assertEqual(workbook["Toma de acción"]["A5"].value, "Prioridad")

    def test_monthly_department_typo_is_normalized(self):
        self.assertEqual(
            friendly_hr_department("ALMACDN MATERIA PRIMA"),
            "ALMACEN MATERIA PRIMA",
        )


if __name__ == "__main__":
    unittest.main()
