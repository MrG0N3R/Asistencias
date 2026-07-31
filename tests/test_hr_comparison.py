from __future__ import annotations

import io
import unittest
from datetime import date, datetime

import pandas as pd
from openpyxl import load_workbook

from hr_dashboard.hr_export import export_hr_incidents_comparison_workbook
from hr_dashboard.hr_incidents import (
    combine_hr_incident_reports,
    compare_hr_incident_periods,
    hr_period_summary,
)


def incident(employee_id, name, incident_type, day, category="Ausencia"):
    return {
        "employee_id": employee_id,
        "name": name,
        "department": "OPERACIONES",
        "source_department": "OPERACIONES",
        "date": day,
        "category": category,
        "incident_type": incident_type,
        "source_description": incident_type,
        "sheet": "Hoja 1",
    }


class HrComparisonTests(unittest.TestCase):
    def setUp(self):
        previous_rows = [
            incident("1", "Ana", "Falta", date(2026, 6, 1)),
            incident("1", "Ana", "Falta", date(2026, 6, 2)),
            incident("2", "Beto", "Permiso", date(2026, 6, 3)),
            incident("4", "Diana", "Incapacidad", date(2026, 6, 4), "Incapacidad"),
            incident("5", "Eva", "Falta", date(2026, 6, 5)),
        ]
        current_rows = [
            incident("1", "Ana", "Falta", date(2026, 7, 1)),
            incident("3", "Carla", "Permiso", date(2026, 7, 2)),
            incident("4", "Diana", "Incapacidad", date(2026, 7, 3), "Incapacidad"),
            incident("5", "Eva", "Falta", date(2026, 7, 4)),
            incident("5", "Eva", "Falta", date(2026, 7, 5)),
        ]
        self.previous = {
            "file_name": "junio.xlsx",
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "report_start": date(2026, 6, 1),
            "report_end": date(2026, 6, 30),
            "incidents": pd.DataFrame(previous_rows),
            "employees_in_report": 4,
        }
        self.current = {
            "file_name": "julio.xlsx",
            "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "report_start": date(2026, 7, 1),
            "report_end": date(2026, 7, 31),
            "incidents": pd.DataFrame(current_rows),
            "employees_in_report": 4,
        }

    def test_reports_are_ordered_and_summarized(self):
        results, incidents = combine_hr_incident_reports([self.current, self.previous])

        self.assertEqual([result["file_name"] for result in results], ["junio.xlsx", "julio.xlsx"])
        self.assertEqual(incidents["report_order"].unique().tolist(), [1, 2])
        summary = hr_period_summary(incidents)
        self.assertEqual(summary["incident_days"].tolist(), [5, 5])
        self.assertEqual(summary["employees"].tolist(), [4, 4])

    def test_all_change_statuses_are_detected(self):
        _, incidents = combine_hr_incident_reports([self.previous, self.current])
        comparison = compare_hr_incident_periods(incidents)
        statuses = {
            (row.employee_id, row.incident_type): row.status
            for row in comparison.itertuples(index=False)
        }

        self.assertEqual(statuses[("1", "Falta")], "Bajó")
        self.assertEqual(statuses[("2", "Permiso")], "Ya no aparece")
        self.assertEqual(statuses[("3", "Permiso")], "Nueva")
        self.assertEqual(statuses[("4", "Incapacidad")], "Se mantiene")
        self.assertEqual(statuses[("5", "Falta")], "Subió")

    def test_empty_current_period_marks_previous_cases_as_missing(self):
        _, incidents = combine_hr_incident_reports([self.previous, self.current])
        previous_only = incidents[incidents["report_order"] == 1]
        comparison = compare_hr_incident_periods(
            previous_only,
            previous_order=1,
            current_order=2,
        )

        self.assertTrue((comparison["status"] == "Ya no aparece").all())
        self.assertTrue((comparison["current_days"] == 0).all())

    def test_comparison_workbook_contains_expected_sheets(self):
        _, incidents = combine_hr_incident_reports([self.previous, self.current])
        workbook_bytes = export_hr_incidents_comparison_workbook(incidents)
        workbook = load_workbook(io.BytesIO(workbook_bytes), read_only=True)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Dashboard comparativo",
                "Cambios por persona",
                "Resumen por periodo",
                "Tendencia por tipo",
                "Tendencia departamento",
                "Detalle consolidado",
            ],
        )


if __name__ == "__main__":
    unittest.main()
