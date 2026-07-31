from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill

from .excel_common import write_dataframe_sheet
from .hr_incidents import (
    compare_hr_incident_periods,
    hr_department_summary,
    hr_employee_summary,
    hr_period_summary,
    summarize_hr_incidents,
)

def export_hr_incidents_workbook(result: dict, incidents: pd.DataFrame) -> bytes:
    summary = summarize_hr_incidents(incidents)
    employee_summary = hr_employee_summary(incidents)
    department_summary = hr_department_summary(incidents)
    type_summary = (
        incidents.groupby(["category", "incident_type"], as_index=False)
        .size()
        .rename(columns={"size": "incident_days"})
        .sort_values("incident_days", ascending=False)
        if not incidents.empty
        else pd.DataFrame(columns=["category", "incident_type", "incident_days"])
    )

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:H2")
    dashboard["A1"] = "Dashboard de incidencias verificadas por RH"
    dashboard["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17365D")
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")

    period_start = result["report_start"]
    period_end = result["report_end"]
    period_text = (
        f"{period_start:%d/%m/%Y} al {period_end:%d/%m/%Y}"
        if period_start and period_end
        else "Periodo no identificado"
    )
    dashboard["A4"] = "Periodo"
    dashboard["B4"] = period_text
    dashboard["E4"] = "Archivo"
    dashboard["F4"] = result["file_name"]
    dashboard["A5"] = "Generado"
    dashboard["B5"] = result["generated_at"]
    for cell in ("A4", "E4", "A5"):
        dashboard[cell].font = Font(bold=True, color="17365D")

    kpis = [
        ("Personas afectadas", summary["employees"], "A7:B7", "A8:B8"),
        ("Días de incidencia", summary["incident_days"], "D7:E7", "D8:E8"),
        ("Ausencias", summary["absences"], "G7:H7", "G8:H8"),
        ("Incapacidades", summary["disabilities"], "A10:B10", "A11:B11"),
        ("Departamentos", summary["departments"], "D10:E10", "D11:E11"),
    ]
    for label, value, label_range, value_range in kpis:
        for row in dashboard[label_range]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="EAF0F6")
        for row in dashboard[value_range]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="EAF0F6")
        dashboard.merge_cells(label_range)
        dashboard.merge_cells(value_range)
        label_cell = dashboard[label_range.split(":")[0]]
        value_cell = dashboard[value_range.split(":")[0]]
        label_cell.value = label
        value_cell.value = value
        label_cell.font = Font(bold=True, color="5B6573")
        value_cell.font = Font(bold=True, size=16, color="17365D")
        label_cell.alignment = Alignment(horizontal="center")
        value_cell.alignment = Alignment(horizontal="center")

    dashboard["A14"] = "Incidencias por tipo"
    dashboard["D14"] = "Departamentos con más días"
    dashboard["G14"] = "Personas con más días"
    for cell in ("A14", "D14", "G14"):
        dashboard[cell].font = Font(bold=True, color="17365D")

    dashboard["A15"] = "Tipo"
    dashboard["B15"] = "Días"
    dashboard["D15"] = "Departamento"
    dashboard["E15"] = "Días"
    dashboard["G15"] = "Empleado"
    dashboard["H15"] = "Días"
    for cell in ("A15", "B15", "D15", "E15", "G15", "H15"):
        dashboard[cell].font = Font(bold=True, color="FFFFFF")
        dashboard[cell].fill = PatternFill("solid", fgColor="0F766E")

    for row_number, record in enumerate(type_summary.head(10).to_dict("records"), start=16):
        dashboard.cell(row_number, 1, record["incident_type"])
        dashboard.cell(row_number, 2, record["incident_days"])
    for row_number, record in enumerate(department_summary.head(10).to_dict("records"), start=16):
        dashboard.cell(row_number, 4, record["department"])
        dashboard.cell(row_number, 5, record["incident_days"])
    for row_number, record in enumerate(employee_summary.head(10).to_dict("records"), start=16):
        dashboard.cell(row_number, 7, record["name"])
        dashboard.cell(row_number, 8, record["incident_days"])

    type_rows = len(type_summary.head(10))
    if type_rows:
        chart = PieChart()
        chart.title = "Distribución por tipo"
        chart.add_data(
            Reference(dashboard, min_col=2, min_row=15, max_row=15 + type_rows),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(dashboard, min_col=1, min_row=16, max_row=15 + type_rows)
        )
        chart.legend = None
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showLeaderLines = True
        chart.height = 7
        chart.width = 12
        dashboard.add_chart(chart, "A28")

    department_rows = len(department_summary.head(10))
    if department_rows:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Días por departamento"
        chart.y_axis.title = "Departamento"
        chart.x_axis.title = "Días"
        chart.add_data(
            Reference(dashboard, min_col=5, min_row=15, max_row=15 + department_rows),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(dashboard, min_col=4, min_row=16, max_row=15 + department_rows)
        )
        chart.legend = None
        chart.height = 7
        chart.width = 14
        dashboard.add_chart(chart, "E28")

    for column, width in {"A": 35, "B": 14, "C": 3, "D": 28, "E": 14, "F": 3, "G": 38, "H": 14}.items():
        dashboard.column_dimensions[column].width = width
    dashboard.freeze_panes = "A7"

    detail = wb.create_sheet("Detalle de incidencias")
    detail_df = incidents.sort_values(["date", "name"]).copy()
    write_dataframe_sheet(
        detail,
        detail_df,
        [
            "No. Trab",
            "Nombre",
            "Departamento homologado",
            "Departamento origen",
            "Fecha",
            "Categoría",
            "Tipo de incidencia",
        ],
        [
            "employee_id",
            "name",
            "department",
            "source_department",
            "date",
            "category",
            "incident_type",
        ],
    )

    employees = wb.create_sheet("Resumen por empleado")
    write_dataframe_sheet(
        employees,
        employee_summary,
        [
            "No. Trab",
            "Nombre",
            "Departamento",
            "Días de incidencia",
            "Ausencias",
            "Incapacidades",
            "Tipos de incidencia",
            "Primera fecha",
            "Última fecha",
        ],
        [
            "employee_id",
            "name",
            "department",
            "incident_days",
            "absences",
            "disabilities",
            "incident_types",
            "first_date",
            "last_date",
        ],
    )

    departments = wb.create_sheet("Resumen por departamento")
    write_dataframe_sheet(
        departments,
        department_summary,
        ["Departamento", "Personas", "Días de incidencia", "Ausencias", "Incapacidades"],
        ["department", "employees", "incident_days", "absences", "disabilities"],
    )

    types = wb.create_sheet("Resumen por tipo")
    write_dataframe_sheet(
        types,
        type_summary,
        ["Categoría", "Tipo de incidencia", "Días"],
        ["category", "incident_type", "incident_days"],
    )

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def export_hr_incidents_comparison_workbook(
    incidents: pd.DataFrame,
    periods: pd.DataFrame | None = None,
    comparison: pd.DataFrame | None = None,
) -> bytes:
    """Genera un Excel con la tendencia y los cambios entre los dos últimos periodos."""
    periods = hr_period_summary(incidents) if periods is None else periods
    comparison = (
        compare_hr_incident_periods(incidents)
        if comparison is None
        else comparison
    )
    current = periods.iloc[-1]
    previous = periods.iloc[-2]

    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard comparativo"
    dashboard.sheet_view.showGridLines = False
    dashboard.merge_cells("A1:J2")
    dashboard["A1"] = "Comparativo de incidencias verificadas por RH"
    dashboard["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor="17365D")
    dashboard["A1"].alignment = Alignment(horizontal="left", vertical="center")

    dashboard["A4"] = "Periodo anterior"
    dashboard["B4"] = previous["report_label"]
    dashboard["F4"] = "Periodo actual"
    dashboard["G4"] = current["report_label"]
    dashboard["A5"] = "Generado"
    dashboard["B5"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    for cell in ("A4", "F4", "A5"):
        dashboard[cell].font = Font(bold=True, color="17365D")

    recurring = int(
        (
            (comparison["previous_days"] > 0)
            & (comparison["current_days"] > 0)
        ).sum()
    )
    improved = int(comparison["status"].isin(["Bajó", "Ya no aparece"]).sum())
    kpis = [
        (
            "Días periodo actual",
            int(current["incident_days"]),
            int(current["incident_days"] - previous["incident_days"]),
        ),
        (
            "Personas periodo actual",
            int(current["employees"]),
            int(current["employees"] - previous["employees"]),
        ),
        ("Incidencias recurrentes", recurring, None),
        ("Bajaron o ya no aparecen", improved, None),
    ]
    for index, (label, value, delta) in enumerate(kpis):
        start_column = 1 + (index * 3)
        label_cell = dashboard.cell(7, start_column, label)
        value_cell = dashboard.cell(8, start_column, value)
        if delta is not None:
            dashboard.cell(9, start_column, f"Variación: {delta:+d}")
        dashboard.merge_cells(
            start_row=7,
            start_column=start_column,
            end_row=7,
            end_column=start_column + 1,
        )
        dashboard.merge_cells(
            start_row=8,
            start_column=start_column,
            end_row=8,
            end_column=start_column + 1,
        )
        dashboard.merge_cells(
            start_row=9,
            start_column=start_column,
            end_row=9,
            end_column=start_column + 1,
        )
        for row_number in (7, 8, 9):
            for column_number in (start_column, start_column + 1):
                cell = dashboard.cell(row_number, column_number)
                cell.fill = PatternFill("solid", fgColor="EAF0F6")
                cell.alignment = Alignment(horizontal="center")
        label_cell.font = Font(bold=True, color="5B6573")
        value_cell.font = Font(bold=True, size=16, color="17365D")

    dashboard["A12"] = "Tendencia por periodo"
    dashboard["A12"].font = Font(bold=True, color="17365D")
    trend_headers = [
        "Periodo",
        "Personas",
        "Días",
        "Ausencias",
        "Incapacidades",
        "Departamentos",
        "Archivo",
    ]
    for column, header in enumerate(trend_headers, start=1):
        cell = dashboard.cell(13, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row_number, row in enumerate(periods.itertuples(index=False), start=14):
        values = [
            row.report_label,
            row.employees,
            row.incident_days,
            row.absences,
            row.disabilities,
            row.departments,
            row.source_file,
        ]
        for column, value in enumerate(values, start=1):
            dashboard.cell(row_number, column, value)

    chart = LineChart()
    chart.title = "Evolución de días de incidencia"
    chart.y_axis.title = "Días"
    chart.x_axis.title = "Periodo"
    chart.add_data(
        Reference(
            dashboard,
            min_col=3,
            min_row=13,
            max_row=13 + len(periods),
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            dashboard,
            min_col=1,
            min_row=14,
            max_row=13 + len(periods),
        )
    )
    chart.legend = None
    chart.height = 7
    chart.width = 16
    dashboard.add_chart(chart, "A21")

    for column, width in {
        "A": 34,
        "B": 18,
        "C": 14,
        "D": 16,
        "E": 18,
        "F": 22,
        "G": 38,
        "H": 4,
        "I": 22,
        "J": 22,
    }.items():
        dashboard.column_dimensions[column].width = width
    dashboard.freeze_panes = "A12"

    changes = wb.create_sheet("Cambios por persona")
    write_dataframe_sheet(
        changes,
        comparison,
        [
            "No. Trab",
            "Nombre",
            "Departamento",
            "Categoría",
            "Tipo de incidencia",
            "Días anteriores",
            "Días actuales",
            "Variación",
            "Variación %",
            "Estado",
        ],
        [
            "employee_id",
            "name",
            "department",
            "category",
            "incident_type",
            "previous_days",
            "current_days",
            "change",
            "change_percent",
            "status",
        ],
    )

    period_sheet = wb.create_sheet("Resumen por periodo")
    write_dataframe_sheet(
        period_sheet,
        periods,
        [
            "Orden",
            "Periodo",
            "Inicio",
            "Fin",
            "Archivo",
            "Personas",
            "Días de incidencia",
            "Ausencias",
            "Incapacidades",
            "Departamentos",
        ],
        [
            "report_order",
            "report_label",
            "report_start",
            "report_end",
            "source_file",
            "employees",
            "incident_days",
            "absences",
            "disabilities",
            "departments",
        ],
    )

    type_trend = (
        incidents.groupby(
            ["report_order", "report_label", "category", "incident_type"],
            as_index=False,
        )
        .size()
        .rename(columns={"size": "incident_days"})
        .sort_values(["report_order", "incident_days"], ascending=[True, False])
    )
    types = wb.create_sheet("Tendencia por tipo")
    write_dataframe_sheet(
        types,
        type_trend,
        ["Orden", "Periodo", "Categoría", "Tipo de incidencia", "Días"],
        ["report_order", "report_label", "category", "incident_type", "incident_days"],
    )

    department_trend = (
        incidents.groupby(["report_order", "report_label", "department"], as_index=False)
        .agg(
            employees=("employee_id", "nunique"),
            incident_days=("date", "size"),
        )
        .sort_values(["report_order", "incident_days"], ascending=[True, False])
    )
    departments = wb.create_sheet("Tendencia departamento")
    write_dataframe_sheet(
        departments,
        department_trend,
        ["Orden", "Periodo", "Departamento", "Personas", "Días"],
        ["report_order", "report_label", "department", "employees", "incident_days"],
    )

    detail = wb.create_sheet("Detalle consolidado")
    detail_df = incidents.sort_values(["report_order", "date", "name"]).copy()
    write_dataframe_sheet(
        detail,
        detail_df,
        [
            "Periodo",
            "Archivo",
            "Fecha",
            "No. Trab",
            "Nombre",
            "Departamento homologado",
            "Departamento origen",
            "Categoría",
            "Tipo de incidencia",
        ],
        [
            "report_label",
            "source_file",
            "date",
            "employee_id",
            "name",
            "department",
            "source_department",
            "category",
            "incident_type",
        ],
    )

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def export_monthly_consolidated_workbook(
    result: dict,
    incidents: pd.DataFrame,
    recommendations: pd.DataFrame,
) -> bytes:
    """Extiende el dashboard mensual con una hoja auditable de toma de acción."""
    workbook_bytes = export_hr_incidents_workbook(result, incidents)
    wb = load_workbook(io.BytesIO(workbook_bytes))
    dashboard = wb["Dashboard"]
    dashboard["A1"] = "Consolidado mensual de incidencias RH"
    dashboard["A13"] = "Consulta la hoja 'Toma de acción' para revisar prioridades y reglas activadas."
    dashboard["A13"].font = Font(bold=True, color="D97706")

    action_sheet = wb.create_sheet("Toma de acción", 1)
    action_sheet.sheet_view.showGridLines = False
    action_sheet.merge_cells("A1:G2")
    action_sheet["A1"] = "Sugerencias de toma de acción"
    action_sheet["A1"].font = Font(bold=True, size=18, color="FFFFFF")
    action_sheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    action_sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    action_sheet.merge_cells("A3:G3")
    action_sheet["A3"] = (
        "Sugerencias operativas generadas con reglas visibles; validar siempre con las políticas internas de RH."
    )
    action_sheet["A3"].font = Font(italic=True, color="5B6573")

    headers = [
        "Prioridad",
        "Hallazgo",
        "Acción sugerida",
        "Alcance",
        "Personas",
        "Días",
        "Regla activada",
    ]
    keys = [
        "priority",
        "finding",
        "suggested_action",
        "scope",
        "affected_employees",
        "incident_days",
        "rule",
    ]
    for column, header in enumerate(headers, start=1):
        cell = action_sheet.cell(5, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    priority_fills = {
        "Alta": "FDE8E7",
        "Media": "FFF3D6",
        "Baja": "E8F5EE",
    }
    for row_number, record in enumerate(recommendations.to_dict("records"), start=6):
        for column, key in enumerate(keys, start=1):
            cell = action_sheet.cell(row_number, column, record.get(key, ""))
            cell.alignment = Alignment(
                horizontal="center" if key in {"priority", "affected_employees", "incident_days"} else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.fill = PatternFill(
                "solid",
                fgColor=priority_fills.get(record.get("priority"), "FFFFFF"),
            )
        action_sheet.row_dimensions[row_number].height = 54

    for column, width in {
        "A": 12,
        "B": 42,
        "C": 58,
        "D": 45,
        "E": 12,
        "F": 10,
        "G": 48,
    }.items():
        action_sheet.column_dimensions[column].width = width
    action_sheet.freeze_panes = "A6"
    action_sheet.auto_filter.ref = f"A5:G{max(5, 5 + len(recommendations))}"

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
