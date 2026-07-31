from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font

from .excel_common import set_widths, write_table

def export_workbook(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dashboard"
    ws["A1"] = "Consolidado de asistencia"
    ws["A1"].font = Font(bold=True, size=16)
    ws.append(["Archivo", result["file_name"]])
    ws.append(["Generado", result["generated_at"]])
    ws.append([])
    for label, value in [
        ("Empleados", result["summary"]["employees"]),
        ("Dias de asistencia", result["summary"]["attendance_days"]),
        ("Retardos", result["summary"]["lates"]),
        ("Incidencias", result["summary"]["incidents"]),
        ("Periodos con falta/bono perdido", result["summary"]["penalties"]),
        ("Salidas faltantes", result["summary"]["missing_exits"]),
        ("Jornadas menores a 8h", result["summary"]["short_shifts"]),
        ("Excepciones de horario activas", len(result.get("schedule_exceptions", []))),
    ]:
        ws.append([label, value])
    ws.append([])
    ws.append(["Regla", "Valor"])
    for key, value in result["rules"].items():
        ws.append([key, value])
    set_widths(ws)

    attendance_keys = [
        "employee_id", "name", "department", "schedule_type", "schedule_source", "building", "date", "period", "shift",
        "scheduled_entry", "scheduled_exit", "entry", "exit", "worked_hours", "marks", "status",
    ]
    ws = wb.create_sheet("Asistencias")
    write_table(
        ws,
        ["No. Trab", "Nombre", "Departamento archivo", "Horario aplicado", "Origen horario", "Edificio", "Fecha operativa", "Periodo", "Turno", "Entrada prog.", "Salida prog.", "Entrada", "Salida", "Horas", "Marcas", "Estatus"],
        result["attendance"],
        attendance_keys,
    )

    ws = wb.create_sheet("Excepciones horario")
    write_table(
        ws,
        ["No. Trab", "Nombre", "Departamento archivo", "Horario asignado"],
        result.get("schedule_exceptions", []),
        ["employee_id", "name", "department", "assigned_schedule"],
    )

    incident_keys = [
        "employee_id", "name", "department", "schedule_type", "schedule_source", "building", "date", "period", "shift",
        "scheduled_entry", "entry", "exit", "worked_hours", "incident_type", "detail", "period_late_count", "bonus_lost",
    ]
    ws = wb.create_sheet("Incidencias")
    write_table(
        ws,
        ["No. Trab", "Nombre", "Departamento archivo", "Horario aplicado", "Origen horario", "Edificio", "Fecha operativa", "Periodo", "Turno", "Entrada prog.", "Entrada", "Salida", "Horas", "Incidencia", "Detalle", "Retardos periodo", "Pierde bono"],
        result["incidents"],
        incident_keys,
    )

    ws = wb.create_sheet("Penalizaciones")
    write_table(
        ws,
        ["No. Trab", "Nombre", "Departamento", "Edificio", "Periodo", "Retardos", "Falta equivalente", "Pierde bono"],
        result["period_penalties"],
        ["employee_id", "name", "department", "building", "period", "late_count", "equivalent_absence", "bonus_lost"],
    )

    ws = wb.create_sheet("Resumen semanal")
    write_table(
        ws,
        ["Periodo", "Dias asistencia", "Retardos", "Jornadas <8h", "Penalizaciones"],
        result["weekly"],
        ["period", "attendance_days", "late_count", "short_shift_count", "penalties"],
    )

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
