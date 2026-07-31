from __future__ import annotations

from datetime import date

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def set_widths(ws):
    for column in ws.columns:
        max_len = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def write_table(ws, headers: list[str], rows: list[dict], keys: list[str]):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5597")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(key, "") for key in keys])
    set_widths(ws)
    ws.freeze_panes = "A2"

def style_excel_header(row) -> None:
    for cell in row:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_dataframe_sheet(ws, dataframe: pd.DataFrame, headers: list[str], keys: list[str]) -> None:
    ws.append(headers)
    style_excel_header(ws[1])
    ws.row_dimensions[1].height = 24
    for record in dataframe.to_dict("records"):
        ws.append([record.get(key, "") for key in keys])
    separator = Side(style="thin", color="D9E3EC")
    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            if isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
            cell.alignment = Alignment(
                horizontal="center" if isinstance(cell.value, (int, float, date)) else "left",
                vertical="center",
            )
            cell.border = Border(bottom=separator)
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    set_widths(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
